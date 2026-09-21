import os
import re
from docx import Document
from docx.shared import Pt

# ============== CONFIGURACIÓN =================
CARPETA_BASE = "."   # raíz del proyecto
SALIDA_DOCX = "codigo_css_por_temas.docx"
# ==============================================

def obtener_numero(nombre):
    match = re.match(r"(\d+)", nombre)
    return int(match.group(1)) if match else 9999

# Crear documento
doc = Document()
doc.add_heading("Código CSS ordenado por temas", level=1)

# Obtener carpetas (solo las que empiezan por número)
carpetas = [
    d for d in os.listdir(CARPETA_BASE)
    if os.path.isdir(d) and re.match(r"\d+_", d)
]

# Ordenar carpetas por número
carpetas.sort(key=obtener_numero)

for carpeta in carpetas:
    ruta_carpeta = os.path.join(CARPETA_BASE, carpeta)

    # Buscar SOLO archivos .css dentro de la carpeta
    archivos_css = [
        f for f in os.listdir(ruta_carpeta)
        if f.lower().endswith(".css")
    ]

    if not archivos_css:
        continue  # saltar carpetas sin CSS

    archivos_css.sort()

    # Título del tema
    doc.add_heading(carpeta.replace("_", " "), level=2)

    for archivo in archivos_css:
        ruta_archivo = os.path.join(ruta_carpeta, archivo)

        # Subtítulo del archivo
        doc.add_heading(f"Archivo CSS: {archivo}", level=3)

        with open(ruta_archivo, "r", encoding="utf-8") as f:
            contenido = f.read()

        # Insertar código EXACTO
        p = doc.add_paragraph()
        run = p.add_run(contenido)
        run.font.name = "Courier New"
        run.font.size = Pt(10)

    doc.add_page_break()

# Guardar documento
doc.save(SALIDA_DOCX)

print("✅ Documento CSS creado correctamente:", SALIDA_DOCX)
